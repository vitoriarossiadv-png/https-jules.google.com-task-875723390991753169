import openpyxl
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.utils import get_column_letter

def create_excel_template():
    wb = openpyxl.Workbook()

    # --- STYLES ---
    # Palette
    COLOR_HEADER_BG = "2C3E50" # Dark Slate Blue
    COLOR_HEADER_TEXT = "FFFFFF"
    COLOR_BORDER = "BDC3C7" # Silver
    COLOR_ACCENT_GOLD = "C0A062" # Champagne Gold
    COLOR_CARD_BG = "F8F9FA" # Off White

    # Fonts
    font_header = Font(name='Arial', size=11, bold=True, color=COLOR_HEADER_TEXT)
    font_body = Font(name='Arial', size=10)
    font_dashboard_title = Font(name='Arial', size=24, bold=True, color=COLOR_HEADER_BG)
    font_metric_label = Font(name='Arial', size=12, color="7F8C8D", bold=True)
    font_metric_value = Font(name='Arial', size=28, bold=True, color=COLOR_HEADER_BG)

    # Borders
    thin_border = Border(left=Side(style='thin', color=COLOR_BORDER),
                         right=Side(style='thin', color=COLOR_BORDER),
                         top=Side(style='thin', color=COLOR_BORDER),
                         bottom=Side(style='thin', color=COLOR_BORDER))

    header_fill = PatternFill(start_color=COLOR_HEADER_BG, end_color=COLOR_HEADER_BG, fill_type="solid")
    card_fill = PatternFill(start_color=COLOR_CARD_BG, end_color=COLOR_CARD_BG, fill_type="solid")

    # ---------------------------------------------------------
    # 1. Config Sheet (Hidden)
    # ---------------------------------------------------------
    ws_config = wb.active
    ws_config.title = "Config"
    ws_config.append(["Complexidade", "Dias"])
    ws_config.append(["Baixa", 4])
    ws_config.append(["Média", 7])
    ws_config.append(["Alta", 10])
    ws_config.sheet_state = 'hidden'

    # ---------------------------------------------------------
    # Helper Function for Sheet Setup
    # ---------------------------------------------------------
    def setup_tracking_sheet(ws_name):
        ws = wb.create_sheet(ws_name)
        headers = [
            "ID", "Cliente", "Tipo de Ação", "Complexidade",
            "Data Envio Doc", "Prazo (Dias)", "Prazo Fatal",
            "Status", "Data Protocolo", "Advogado Responsável"
        ]
        ws.append(headers)

        # Apply Header Styles
        for col_idx, cell in enumerate(ws[1], 1):
            cell.fill = header_fill
            cell.font = font_header
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = thin_border

        # Set Column Widths
        col_widths = [8, 35, 25, 15, 18, 12, 18, 18, 18, 25]
        for i, width in enumerate(col_widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = width

        # Add Data Validation (Complexidade & Status)
        dv_complexity = DataValidation(type="list", formula1='"Baixa,Média,Alta"', showDropDown=True)
        ws.add_data_validation(dv_complexity)
        dv_complexity.add("D2:D1000")

        dv_status = DataValidation(type="list", formula1='"Pendente,Em Andamento,Finalizada"', showDropDown=True)
        ws.add_data_validation(dv_status)
        dv_status.add("H2:H1000")

        # Add Formulas & Default Styling for Rows
        for row in range(2, 101):
            for col in range(1, 11):
                cell = ws.cell(row=row, column=col)
                cell.font = font_body
                cell.border = thin_border
                cell.alignment = Alignment(vertical="center")

                if col not in [2, 3, 10]: # Center align except text fields
                    cell.alignment = Alignment(horizontal="center", vertical="center")

            # Formulas
            ws[f"F{row}"] = f'=IFERROR(VLOOKUP(D{row}, Config!A:B, 2, FALSE), "")'
            ws[f"G{row}"] = f'=IF(E{row}="", "", WORKDAY(E{row}, F{row}))'

        # Conditional Formatting
        green_fill = PatternFill(start_color='C6EFCE', end_color='C6EFCE', fill_type='solid')
        yellow_fill = PatternFill(start_color='FFEB9C', end_color='FFEB9C', fill_type='solid')
        red_fill = PatternFill(start_color='FFC7CE', end_color='FFC7CE', fill_type='solid')

        # Complexity
        ws.conditional_formatting.add('D2:D1000', CellIsRule(operator='equal', formula=['"Baixa"'], fill=green_fill))
        ws.conditional_formatting.add('D2:D1000', CellIsRule(operator='equal', formula=['"Média"'], fill=yellow_fill))
        ws.conditional_formatting.add('D2:D1000', CellIsRule(operator='equal', formula=['"Alta"'], fill=red_fill))

        # Status
        ws.conditional_formatting.add('H2:H1000', CellIsRule(operator='equal', formula=['"Finalizada"'], fill=green_fill))

        return ws

    # ---------------------------------------------------------
    # 2. Entrada (Input Sheet)
    # ---------------------------------------------------------
    setup_tracking_sheet("Entrada")

    # ---------------------------------------------------------
    # 3. Monthly Sheets
    # ---------------------------------------------------------
    months = ["Jan", "Fev", "Mar", "Abr", "Mai", "Jun",
              "Jul", "Ago", "Set", "Out", "Nov", "Dez"]
    for month in months:
        setup_tracking_sheet(month)

    # ---------------------------------------------------------
    # 4. Dashboard (Redesigned)
    # ---------------------------------------------------------
    ws_dash = wb.create_sheet("Dashboard", 0)
    ws_dash.sheet_view.showGridLines = False

    # Title
    ws_dash['B2'] = "DASHBOARD - GESTÃO DE PROTOCOLOS"
    ws_dash['B2'].font = font_dashboard_title
    ws_dash['B2'].alignment = Alignment(horizontal="left")

    def create_card(start_row, start_col, title, formula, number_format="0"):
        # Style Title
        title_cell = ws_dash.cell(row=start_row, column=start_col)
        title_cell.value = title
        title_cell.font = font_metric_label
        title_cell.alignment = Alignment(horizontal="center", vertical="center")
        title_cell.fill = card_fill

        # Style Value
        value_cell = ws_dash.cell(row=start_row+1, column=start_col)
        value_cell.value = formula
        value_cell.font = font_metric_value
        value_cell.alignment = Alignment(horizontal="center", vertical="center")
        value_cell.number_format = number_format

        # Merge
        ws_dash.merge_cells(start_row=start_row, start_column=start_col, end_row=start_row, end_column=start_col+2)
        ws_dash.merge_cells(start_row=start_row+1, start_column=start_col, end_row=start_row+2, end_column=start_col+2)

        # Borders
        # Draw border around the whole block (3x3)
        for r in range(start_row, start_row+3):
            for c in range(start_col, start_col+3):
                cell = ws_dash.cell(row=r, column=c)
                # Apply outer border logic
                sides = {'left': None, 'right': None, 'top': None, 'bottom': None}

                if c == start_col: sides['left'] = Side(style='thin', color=COLOR_BORDER)
                if c == start_col+2: sides['right'] = Side(style='thin', color=COLOR_BORDER)
                if r == start_row: sides['top'] = Side(style='thin', color=COLOR_BORDER)
                if r == start_row+2: sides['bottom'] = Side(style='thin', color=COLOR_BORDER)

                cell.border = Border(**sides)

                # Internal divider
                if r == start_row:
                    cell.border = Border(bottom=Side(style='thin', color=COLOR_ACCENT_GOLD), **{k:v for k,v in sides.items() if k != 'bottom'})

    # Metrics Layout
    total_finalized_sum = "+".join([f"(COUNTA({m}!A:A)-1)" for m in months])
    formula_total = f'=COUNTA(Entrada!A:A)-1 + {total_finalized_sum}'

    create_card(5, 2, "TOTAL PROTOCOLOS", formula_total) # B5

    formula_finalized = f'={total_finalized_sum}'
    create_card(5, 6, "FINALIZADAS", formula_finalized) # F5

    formula_meta = f'=IFERROR({formula_finalized} / {formula_total}, 0)'
    create_card(5, 10, "% META BATIDA", formula_meta, "0%") # J5

    # Status Indicator
    ws_dash['B10'] = "STATUS DA META"
    ws_dash['B10'].font = font_metric_label

    # Logic: <0.8 = Red, 0.8-0.99 = Blue, 1.0 = Green
    ws_dash['B11'] = f'=IF({formula_meta}=1, "SUPERMETA 🚀", IF({formula_meta}>=0.8, "META BATIDA ✅", "META NÃO BATIDA ⚠️"))'
    ws_dash['B11'].font = Font(name='Arial', size=20, bold=True)
    ws_dash.merge_cells('B11:E12')
    ws_dash['B11'].alignment = Alignment(horizontal="left", vertical="center")

    # Conditional Formatting for Status Text
    # Use cell references to J6 (where % meta is) - wait, J6 is merged?
    # J5 is title, J6 is value (merged J6:L7). referencing J6 works.
    # Note: openpyxl formulas in rules must be strings.

    # We need to refer to the value cell of the % Meta card.
    # create_card at (5, 10) -> Title at (5,10), Value at (6,10) i.e., J6.

    # Green
    ws_dash.conditional_formatting.add('B11', FormulaRule(formula=['J6=1'], font=Font(color="27AE60", size=20, bold=True)))
    # Blue/Yellow
    ws_dash.conditional_formatting.add('B11', FormulaRule(formula=['AND(J6>=0.8, J6<1)'], font=Font(color="2980B9", size=20, bold=True)))
    # Red
    ws_dash.conditional_formatting.add('B11', FormulaRule(formula=['J6<0.8'], font=Font(color="C0392B", size=20, bold=True)))

    # Adjust widths
    ws_dash.column_dimensions['A'].width = 2
    for col in ['B', 'C', 'D', 'F', 'G', 'H', 'J', 'K', 'L']:
        ws_dash.column_dimensions[col].width = 15
    ws_dash.column_dimensions['E'].width = 3
    ws_dash.column_dimensions['I'].width = 3

    wb.save("legal_control.xlsx")
    print("Excel file 'legal_control.xlsx' created successfully with Enhanced Styles.")

if __name__ == "__main__":
    create_excel_template()
