import openpyxl

def verify_excel():
    try:
        wb = openpyxl.load_workbook("legal_control.xlsx")
        print("Sheets found:", wb.sheetnames)

        expected_sheets = ["Config", "Dashboard", "Entrada", "Jan", "Fev", "Mar", "Abr", "Mai", "Jun", "Jul", "Ago", "Set", "Out", "Nov", "Dez"]

        missing_sheets = [s for s in expected_sheets if s not in wb.sheetnames]
        if missing_sheets:
            print("ERROR: Missing sheets:", missing_sheets)
        else:
            print("All expected sheets are present.")

        ws_entrada = wb["Entrada"]
        headers = [cell.value for cell in ws_entrada[1]]
        print("Headers in 'Entrada':", headers)

        expected_headers = ["ID", "Cliente", "Tipo de Ação", "Complexidade", "Data Envio Doc", "Prazo (Dias)", "Prazo Fatal", "Status", "Data Protocolo", "Advogado Responsável"]

        if headers == expected_headers:
            print("Headers match expected structure.")
        else:
            print("ERROR: Headers do not match expected structure.")
            print("Expected:", expected_headers)
            print("Found:", headers)

    except Exception as e:
        print(f"Verification failed: {e}")

if __name__ == "__main__":
    verify_excel()
